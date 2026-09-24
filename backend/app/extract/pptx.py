"""Extractor for .pptx files (python-pptx).

Slide titles become heading blocks (level 1); other text-frame shapes and
speaker notes become paragraph blocks; tables serialise to Markdown. No
list-style detection — PPTX bullets aren't structurally distinguishable the
way DOCX list styles are. `anchor` is a sequential block index.
"""

from pathlib import Path

from pptx import Presentation
from pptx.table import Table as PptxTable

from app.extract.base import Block


def _table_to_markdown(table: PptxTable) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    header, *body = rows
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def extract(path: Path) -> list[Block]:
    presentation = Presentation(str(path))
    blocks: list[Block] = []
    anchor = 0

    for slide in presentation.slides:
        title_shape = slide.shapes.title
        title_shape_id = title_shape.shape_id if title_shape is not None else None

        if title_shape is not None and title_shape.has_text_frame:
            title_text = title_shape.text_frame.text.strip()
            if title_text:
                blocks.append(Block(type="heading", text=title_text, level=1, anchor=anchor))
                anchor += 1

        for shape in slide.shapes:
            # slide.shapes.title returns a fresh wrapper object each access,
            # so compare by shape_id rather than identity.
            if title_shape_id is not None and shape.shape_id == title_shape_id:
                continue

            if shape.has_table:
                table_text = _table_to_markdown(shape.table)
                if table_text:
                    blocks.append(Block(type="table", text=table_text, level=None, anchor=anchor))
                    anchor += 1
                continue

            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        blocks.append(
                            Block(type="paragraph", text=text, level=None, anchor=anchor)
                        )
                        anchor += 1

        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                blocks.append(Block(type="paragraph", text=notes_text, level=None, anchor=anchor))
                anchor += 1

    return blocks
